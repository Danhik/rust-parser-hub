import os
import sys
import json
import time
import threading
import subprocess
import uuid
import shutil
from pathlib import Path
from typing import Dict, Any, Optional, List

from flask import Flask, jsonify, request, send_from_directory, send_file
from openpyxl import Workbook, load_workbook

# ---------------------------------------------------------------------------
# Paths  (project is now self-contained inside parser_hub/)
# ---------------------------------------------------------------------------

BASE_DIR           = Path(__file__).parent
PARSERS_DIR        = BASE_DIR / "parsers"          # parser scripts
DATA_DIR           = BASE_DIR / "data"             # excel outputs
SDA_DIR            = BASE_DIR / "sda_profiles"     # SDA .maFile folders/files
STATIC_DIR         = BASE_DIR / "static"
CONFIG_FILE        = BASE_DIR / "config.json"
EXTERNAL_DATA_DIR  = BASE_DIR / "external_data"   # external CSV catalogs (avan, lisskins…)
CHROME_DIR         = BASE_DIR / "chrome_profiles"  # chrome profile folders


PARSERS_DIR.mkdir(exist_ok=True)
CHROME_DIR.mkdir(exist_ok=True)
DATA_DIR.mkdir(exist_ok=True)
SDA_DIR.mkdir(exist_ok=True)
EXTERNAL_DATA_DIR.mkdir(exist_ok=True)
# ---------------------------------------------------------------------------
# Config
# ---------------------------------------------------------------------------

DEFAULT_CONFIG = {
    "proxies_file": "proxies.txt",
    "names_file":   "names.txt",
    "default_appid": "252490",
    "default_game_id": "252490",
    "outputs": {
        "steam_market":        "steam_market.xlsx",
        "steam_priceoverview": "steam_priceoverview.xlsx",
        "swapgg_site":         "swapgg_site.xlsx",
        "swapgg_user":         "swapgg_user.xlsx",
        "skinswap_site":       "skinswap_site.xlsx",
        "skinswap_user":       "skinswap_user.xlsx",
        "tradeit_site":        "tradeit_site.xlsx",
        "tradeit_user":        "tradeit_user.xlsx",
    },
    "workers": {
        "steam_market": 4,
        "steam_priceoverview": 20,
        "tradeit_site": 5,
        "skinswap_user": 3,
    },
}


def _deep_merge(base: dict, override: dict):
    for k, v in override.items():
        if k in base and isinstance(base[k], dict) and isinstance(v, dict):
            _deep_merge(base[k], v)
        else:
            base[k] = v


def load_config() -> dict:
    cfg = json.loads(json.dumps(DEFAULT_CONFIG))
    if CONFIG_FILE.exists():
        try:
            saved = json.loads(CONFIG_FILE.read_text(encoding="utf-8"))
            _deep_merge(cfg, saved)
        except Exception:
            pass
    return cfg


def save_config(cfg: dict):
    CONFIG_FILE.write_text(json.dumps(cfg, ensure_ascii=False, indent=2), encoding="utf-8")

# ---------------------------------------------------------------------------
# Flask app
# ---------------------------------------------------------------------------

app = Flask(__name__, static_folder=str(STATIC_DIR), static_url_path="/static")

# ---------------------------------------------------------------------------
# Task manager
# ---------------------------------------------------------------------------

tasks: Dict[str, Dict[str, Any]] = {}
tasks_lock = threading.Lock()


def _run_task(task_id: str, cmd: List[str], cwd: str):
    with tasks_lock:
        tasks[task_id].update({"status": "running", "started_at": time.time()})

    log_lines: List[str] = []
    env = os.environ.copy()
    env["PYTHONIOENCODING"] = "utf-8"
    env["PYTHONUTF8"] = "1"
    try:
        proc = subprocess.Popen(
            cmd, cwd=cwd,
            stdout=subprocess.PIPE, stderr=subprocess.STDOUT,
            text=True, encoding="utf-8", errors="replace",
            env=env,
        )
        with tasks_lock:
            tasks[task_id]["pid"] = proc.pid

        for line in proc.stdout:
            log_lines.append(line.rstrip("\n"))
            with tasks_lock:
                tasks[task_id]["log"] = log_lines[-300:]

        proc.wait()
        status = "done" if proc.returncode == 0 else "error"
        with tasks_lock:
            tasks[task_id].update({"status": status, "returncode": proc.returncode,
                                    "log": log_lines[-300:]})
    except Exception as e:
        with tasks_lock:
            tasks[task_id].update({"status": "error", "error": str(e), "log": log_lines[-300:]})

# ---------------------------------------------------------------------------
# Parser definitions
# ---------------------------------------------------------------------------

PARSERS = {
    # Site inventory  — что площадка ОТДАЁТ (цена продажи площадки)
    "steam_market":        "steam_market_parser.py",
    "steam_priceoverview": "steam_price_overview_parser.py",
    "swapgg_site":         "swapgg_selenium_network_parser.py",
    "skinswap_site":       "skinswap_browser_parser.py",
    "tradeit_site":        "tradeit_parser.py",
    # User inventory  — что площадка ПРИНИМАЕТ от пользователя (цена покупки площадки)
    "swapgg_user":         "swapgg_user_inventory_parser.py",
    "skinswap_user":       "skinswap_user_inventory_parser.py",
    "tradeit_user":        "tradeit_user_inventory_parser.py",
}

# Indicates which parsers are "user inventory" (platform ACCEPTS)
USER_INVENTORY_PARSERS = {"swapgg_user", "skinswap_user", "tradeit_user"}

# ---------------------------------------------------------------------------
# Price parsing
# ---------------------------------------------------------------------------

def _parse_price(val) -> Optional[float]:
    if val is None or val == "":
        return None
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).strip().replace("$", "").replace(" ", "").replace(",", ".")
    try:
        return float(s)
    except Exception:
        return None


def _col_index(headers: List[str], keywords: List[str]) -> Optional[int]:
    lower = [str(h).lower() for h in headers]
    for kw in keywords:
        for i, h in enumerate(lower):
            if kw in h:
                return i
    return None


def read_excel_items(filepath: str) -> List[Dict[str, Any]]:
    try:
        wb  = load_workbook(filepath, read_only=True, data_only=True)
        ws  = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if len(rows) < 2:
            return []
        headers   = [str(h) if h is not None else "" for h in rows[0]]
        name_col  = _col_index(headers, ["name"])
        price_col = _col_index(headers, ["price (trade)", "price_usd", "price"])
        if name_col is None or price_col is None:
            return []
        items = []
        for row in rows[1:]:
            if not row:
                continue
            name  = str(row[name_col]).strip() if row[name_col] is not None else ""
            price = _parse_price(row[price_col]) if len(row) > price_col else None
            if name and price is not None and price > 0:
                items.append({"name": name, "price": price})
        wb.close()
        return items
    except Exception:
        return []


def read_csv_items(filepath: str) -> List[Dict[str, Any]]:
    """Read semicolon-separated CSV (Avan / LisSkins format).

    Supports:
    - Optional header row (auto-detected: if cell 2 in row 1 is not numeric → header)
    - Quoted field values (strips leading/trailing quotes)
    - Comma as decimal separator  (e.g. '12,34' → 12.34)
    """
    import csv
    items: List[Dict[str, Any]] = []
    try:
        with open(filepath, encoding="utf-8", errors="replace", newline="") as fh:
            reader = csv.reader(fh, delimiter=";")
            raw_rows = [row for row in reader if row]
        if not raw_rows:
            return []

        def _strip(s: str) -> str:
            return s.strip().strip('"').strip("'").strip()

        # Auto-detect header: if 2nd cell of row 0 is non-numeric it's a header
        first_price_str = _strip(raw_rows[0][1]) if len(raw_rows[0]) > 1 else ""
        has_header = _parse_price(first_price_str.replace(",", ".")) is None
        data_rows = raw_rows[1:] if has_header else raw_rows

        for row in data_rows:
            if len(row) < 2:
                continue
            name  = _strip(row[0])
            price = _parse_price(_strip(row[1]).replace(",", "."))
            if name and price is not None and price > 0:
                items.append({"name": name, "price": price})
    except Exception:
        pass
    return items


def read_items(filepath: str) -> List[Dict[str, Any]]:
    """Unified reader: dispatches to Excel or CSV reader based on file extension."""
    ext = Path(filepath).suffix.lower()
    if ext == ".csv":
        return read_csv_items(filepath)
    return read_excel_items(filepath)

# ---------------------------------------------------------------------------
# Static
# ---------------------------------------------------------------------------

@app.route("/")
def index():
    return send_from_directory(str(STATIC_DIR), "index.html")

# ---------------------------------------------------------------------------
# Config API
# ---------------------------------------------------------------------------

@app.route("/api/config", methods=["GET"])
def get_config():
    return jsonify(load_config())


@app.route("/api/config", methods=["POST"])
def set_config():
    cfg = load_config()
    _deep_merge(cfg, request.json or {})
    save_config(cfg)
    return jsonify({"ok": True, "config": cfg})

# ---------------------------------------------------------------------------
# Environment scan
# ---------------------------------------------------------------------------

@app.route("/api/scan", methods=["GET"])
def scan_environment():
    result = {
        "chrome_profiles": [],
        "session_profiles": {},
        "sda_profiles":    [],
        "proxy_files":     [],
    }

    # Chrome profiles from dedicated folder
    for item in CHROME_DIR.iterdir():
        if item.is_dir():
            result["chrome_profiles"].append({"name": item.name, "path": str(item)})

    # Also detect chrome profile dirs in parsers/ and base
    for search_dir in [PARSERS_DIR, BASE_DIR]:
        for item in search_dir.iterdir():
            if not item.is_dir():
                continue
            if ("chrome" in item.name.lower() or "profile" in item.name.lower()) and \
               item.name not in ("parsers", "data", "static", "chrome_profiles", "__pycache__"):
                entry = {"name": item.name, "path": str(item)}
                if entry not in result["chrome_profiles"]:
                    result["chrome_profiles"].append(entry)

    # Session JSON files  (stored in parsers/ when scripts run)
    for f in PARSERS_DIR.glob("tradeit_session_*.json"):
        name = f.name.replace("tradeit_session_", "").replace(".json", "")
        result["session_profiles"].setdefault("tradeit", []).append(name)
    for f in PARSERS_DIR.glob("skinswap_session_*.json"):
        name = f.name.replace("skinswap_session_", "").replace(".json", "")
        result["session_profiles"].setdefault("skinswap", []).append(name)

    # SDA .maFile — scan parsers/ and sda_profiles/
    for search_dir in [PARSERS_DIR, SDA_DIR]:
        if search_dir.exists():
            for f in search_dir.glob("*.maFile"):
                if f.stem not in result["sda_profiles"]:
                    result["sda_profiles"].append(f.stem)

    # Proxy files in parsers/ dir
    for f in PARSERS_DIR.glob("*.txt"):
        try:
            first = f.read_text(encoding="utf-8", errors="ignore").split("\n")[0].strip()
            parts = first.split(":")
            if len(parts) >= 2 and parts[1].isdigit():
                result["proxy_files"].append(f.name)
        except Exception:
            pass
    if not result["proxy_files"]:
        result["proxy_files"] = ["proxies.txt"]

    return jsonify(result)

# ---------------------------------------------------------------------------
# Chrome profile management
# ---------------------------------------------------------------------------

@app.route("/api/chrome-profiles", methods=["GET"])
def list_chrome_profiles():
    profiles = []
    for item in CHROME_DIR.iterdir():
        if item.is_dir():
            profiles.append({"name": item.name, "path": str(item)})
    return jsonify(profiles)


@app.route("/api/create-chrome-profile", methods=["POST"])
def create_chrome_profile():
    body = request.json or {}
    name = body.get("name", "").strip()
    if not name:
        return jsonify({"error": "Name is required"}), 400
    # Sanitize name
    name = "".join(c for c in name if c.isalnum() or c in "-_")
    if not name:
        return jsonify({"error": "Invalid name (use letters, digits, - _)"}), 400

    profile_path = CHROME_DIR / name
    if profile_path.exists():
        return jsonify({"error": f"Profile '{name}' already exists", "path": str(profile_path)}), 409

    profile_path.mkdir(parents=True)
    return jsonify({"ok": True, "name": name, "path": str(profile_path)})


@app.route("/api/delete-chrome-profile", methods=["POST"])
def delete_chrome_profile():
    body = request.json or {}
    name = body.get("name", "").strip()
    if not name:
        return jsonify({"error": "Name is required"}), 400
    profile_path = CHROME_DIR / name
    if not profile_path.exists():
        return jsonify({"error": "Profile not found"}), 404
    try:
        shutil.rmtree(profile_path)
        return jsonify({"ok": True})
    except Exception as e:
        return jsonify({"error": str(e)}), 500

# ---------------------------------------------------------------------------
# Parser run
# ---------------------------------------------------------------------------

@app.route("/api/parsers", methods=["GET"])
def list_parsers():
    return jsonify([
        {"id": pid, "script": script, "exists": (PARSERS_DIR / script).exists(),
         "type": "user_inventory" if pid in USER_INVENTORY_PARSERS else "site_inventory"}
        for pid, script in PARSERS.items()
    ])


@app.route("/api/run-parser", methods=["POST"])
def run_parser():
    body      = request.json or {}
    parser_id = body.get("parser_id", "")
    if parser_id not in PARSERS:
        return jsonify({"error": f"Unknown parser: {parser_id}"}), 400

    script      = PARSERS[parser_id]
    script_path = PARSERS_DIR / script
    if not script_path.exists():
        return jsonify({"error": f"Script not found: {script}"}), 404

    out_filename = body.get("out", f"{parser_id}_{int(time.time())}.xlsx")
    out_path = str(DATA_DIR / out_filename) if not os.path.isabs(out_filename) else out_filename

    extra = body.get("args", {})
    cmd   = [sys.executable, "-u", str(script_path), "--out", out_path]

    arg_map = {
        "proxies":       "--proxies",
        "appid":         "--appid",
        "workers":       "--workers",
        "proxy_passes":  "--proxy-passes",
        "game_id":       "--game-id",
        "user_profile":  "--user-profile",
        "user_data_dir": "--user-data-dir",
        "max_items":     "--max-items",
        "names":         "--names",
        "sort":          "--sort",
        "timeout":       "--timeout",
        "total_count":   "--total-count",
    }
    bool_args = ["verbose", "accepted_only", "append", "repair", "headless"]

    for key, flag in arg_map.items():
        val = extra.get(key)
        if val is not None and str(val).strip():
            cmd.extend([flag, str(val)])

    for key in bool_args:
        if extra.get(key):
            cmd.append("--" + key.replace("_", "-"))

    task_id = str(uuid.uuid4())
    with tasks_lock:
        tasks[task_id] = {
            "id": task_id, "parser_id": parser_id, "script": script,
            "status": "pending", "log": [], "out": out_path,
            "cmd": " ".join(cmd), "created_at": time.time(),
        }

    threading.Thread(target=_run_task, args=(task_id, cmd, str(PARSERS_DIR)), daemon=True).start()
    return jsonify({"task_id": task_id, "out": out_path})


@app.route("/api/task-status/<task_id>")
def task_status(task_id):
    with tasks_lock:
        task = tasks.get(task_id)
    if not task:
        return jsonify({"error": "Not found"}), 404
    return jsonify(task)


@app.route("/api/tasks")
def list_tasks():
    with tasks_lock:
        result = list(tasks.values())
    result.sort(key=lambda t: t.get("created_at", 0), reverse=True)
    return jsonify(result[:50])


@app.route("/api/kill-task/<task_id>", methods=["POST"])
@app.route("/api/cancel-task/<task_id>", methods=["POST"])
def kill_task(task_id):
    with tasks_lock:
        task = tasks.get(task_id)
    if not task:
        return jsonify({"error": "Not found"}), 404
    pid = task.get("pid")
    if pid:
        try:
            subprocess.call(["taskkill", "/PID", str(pid), "/F", "/T"],
                            stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        except Exception:
            pass
    with tasks_lock:
        if task_id in tasks:
            tasks[task_id]["status"] = "cancelled"
    return jsonify({"ok": True})

# ---------------------------------------------------------------------------
# Files
# ---------------------------------------------------------------------------

@app.route("/api/list-files")
def list_files():
    files = []
    for f in DATA_DIR.iterdir():
        if f.suffix.lower() == ".xlsx":
            files.append({"name": f.name, "path": str(f),
                           "size": f.stat().st_size, "modified": f.stat().st_mtime})
    files.sort(key=lambda x: x["modified"], reverse=True)
    return jsonify(files)



@app.route("/api/list-external-files")
def list_external_files():
    """List CSV files in external_data/ directory."""
    files = []
    if EXTERNAL_DATA_DIR.exists():
        for f in EXTERNAL_DATA_DIR.iterdir():
            if f.suffix.lower() == ".csv":
                files.append({
                    "name": f.name,
                    "path": str(f),
                    "size": f.stat().st_size,
                    "modified": f.stat().st_mtime,
                })
    files.sort(key=lambda x: x["name"])
    return jsonify(files)


@app.route("/api/download/<filename>")
def download_file(filename):
    fp = DATA_DIR / filename
    if not fp.exists():
        return jsonify({"error": "Not found"}), 404
    return send_file(str(fp), as_attachment=True)

# ---------------------------------------------------------------------------
# Links builder  —  ПЕРЕРАБОТАН
#
# Логика:
#   source_file  = инвентарь сайта  (площадка ОТДАЁТ, цена = стоимость покупки у площадки)
#   dest_file    = инвентарь пользователя (площадка ПРИНИМАЕТ, цена = сколько площадка платит тебе)
#
#   Прибыль = dest_price - source_price
#   Profit % = profit / source_price * 100
#
# Опционально: коэффициенты для учёта доп. расходов (Steam market fee и т.д.)
# ---------------------------------------------------------------------------

@app.route("/api/build-links", methods=["POST"])
def build_links():
    body = request.json or {}

    def _resolve(name_or_path: str) -> Optional[str]:
        if not name_or_path:
            return None
        p = Path(name_or_path)
        if p.is_absolute() and p.exists():
            return str(p)
        c = DATA_DIR / name_or_path
        if c.exists():
            return str(c)
        e = EXTERNAL_DATA_DIR / name_or_path
        if e.exists():
            return str(e)
        return None

    source_file = _resolve(body.get("source_file", ""))
    dest_file   = _resolve(body.get("dest_file",   ""))

    if not source_file:
        return jsonify({"error": "source_file не найден"}), 400
    if not dest_file:
        return jsonify({"error": "dest_file не найден"}), 400

    source_coeff = float(body.get("source_coeff", 1.0))
    dest_coeff   = float(body.get("dest_coeff",   1.0))
    min_profit   = float(body.get("min_profit", 0))
    min_roi      = float(body.get("min_roi",    0))

    source_items = read_items(source_file)
    dest_items   = read_items(dest_file)

    if not source_items:
        return jsonify({"error": "Нет данных в source_file (нужны колонки Name и Price)"}), 400
    if not dest_items:
        return jsonify({"error": "Нет данных в dest_file (нужны колонки Name и Price)"}), 400

    # Build lookup: name -> best (highest) accept price on destination
    # (We want highest because that's best for us as sellers)
    dest_lookup: Dict[str, float] = {}
    for item in dest_items:
        name  = item["name"]
        price = item["price"]
        if name not in dest_lookup or price > dest_lookup[name]:
            dest_lookup[name] = price

    results = []
    for item in source_items:
        name       = item["name"]
        buy_raw    = item["price"]               # что площадка просит за предмет
        accept_raw = dest_lookup.get(name)
        if accept_raw is None:
            continue

        buy_eff    = buy_raw    * source_coeff   # реальная стоимость с учётом коэфф.
        accept_eff = accept_raw * dest_coeff     # реальная выручка с учётом коэфф.
        profit     = accept_eff - buy_eff
        roi        = (profit / buy_eff * 100) if buy_eff > 0 else 0

        if profit < min_profit:
            continue
        if roi < min_roi:
            continue

        results.append({
            "name":        name,
            "source_price":  round(buy_raw,    2),  # цена площадки (отдаёт)
            "dest_price":    round(accept_raw,  2),  # цена принятия
            "source_eff":  round(buy_eff,    2),
            "dest_eff":    round(accept_eff, 2),
            "profit":      round(profit,     2),
            "roi":         round(roi,        1),
        })

    results.sort(key=lambda x: x["profit"], reverse=True)

    out_file = body.get("out_file", f"links_{int(time.time())}.xlsx")
    out_path = DATA_DIR / out_file
    _save_links_excel(results, str(out_path), source_file, dest_file, source_coeff, dest_coeff)

    return jsonify({
        "count":        len(results),
        "items":        results[:500],
        "out_file":     out_file,
        "source_total": len(source_items),
        "dest_total":   len(dest_items),
        "matched":      len(results),
    })


def _save_links_excel(results, out_path, source_file, dest_file, sc, dc):
    wb = Workbook()
    ws = wb.active
    ws.title = "Связки"
    ws.append(["Откуда (отдаёт):",  Path(source_file).name, f"коэфф. x{sc}"])
    ws.append(["Куда  (принимает):", Path(dest_file).name,   f"коэфф. x{dc}"])
    ws.append([])
    ws.append(["Название предмета",
               "Цена продажи (площадка отдаёт)",
               "Цена принятия (площадка берёт)",
               "Покупка с коэфф.",
               "Принятие с коэфф.",
               "Прибыль",
               "ROI %"])
    for r in results:
        ws.append([r["name"], r["source_price"], r["dest_price"],
                   r["source_eff"], r["dest_eff"], r["profit"], r["roi"]])
    ws.column_dimensions["A"].width = 50
    for c in "BCDEFG":
        ws.column_dimensions[c].width = 18
    try:
        wb.save(out_path)
    except Exception as e:
        print(f"[links] save error: {e}")

# ---------------------------------------------------------------------------
# Profiles
# ---------------------------------------------------------------------------

@app.route("/api/profiles")
def list_profiles():
    profiles = []
    for f in PARSERS_DIR.glob("tradeit_session_*.json"):
        profiles.append({"platform": "tradeit",
                          "name": f.name.replace("tradeit_session_", "").replace(".json", ""),
                          "file": f.name})
    for f in PARSERS_DIR.glob("skinswap_session_*.json"):
        profiles.append({"platform": "skinswap",
                          "name": f.name.replace("skinswap_session_", "").replace(".json", ""),
                          "file": f.name})
    return jsonify(profiles)

# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

if __name__ == "__main__":
    print("Parser Hub starting...")
    print(f"  Scripts  : {PARSERS_DIR}")
    print(f"  Data     : {DATA_DIR}")
    print(f"  Profiles : {CHROME_DIR}")
    print(f"  Config   : {CONFIG_FILE}")
    print(f"  >> http://localhost:5050")
    app.run(host="0.0.0.0", port=5050, debug=False)
